"""Ingestion for the FFU Analyzer.

Reads every document in the FFU package one page at a time and stores it in
SQLite, together with the two things the plain text does not tell you:

  * which documents have been replaced by a newer revision, and
  * which passages inside a revised document were struck out or highlighted.

KFU 1 states the convention used in this package: changes are marked in yellow
and struck-through codes no longer apply.  Text extraction cannot see either,
so a naive reader treats deleted requirements as if they still applied.
"""

import logging
import re
import sqlite3
import unicodedata
import zipfile
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

SKIP_FILES = {".DS_Store", "extraction.json", "plot.log", "ref_patterns.json"}
SUPPORTED = {".pdf", ".xlsx", ".xlsm"}
CHUNK_CHARS = 1400

# --------------------------------------------------------------------------
# filenames
# --------------------------------------------------------------------------


def fix_name(name: str) -> str:
    """Repair Mac-zip filenames that were extracted on Windows.

    The zip stores "ö" decomposed (o + U+0308).  Extracted with the cp437
    fallback those bytes turn into box characters, so "Ritningsförteckning"
    reaches us as "Ritningsfo╠êrteckning".  Encoding back to cp437 and
    decoding as UTF-8 undoes it; NFC then composes the accents.
    """
    try:
        name = name.encode("cp437").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        pass
    return unicodedata.normalize("NFC", name)


# --------------------------------------------------------------------------
# document metadata
# --------------------------------------------------------------------------

REV_RE = re.compile(r"\brev\.?\s*(\d{4}-\d{2}-\d{2})", re.I)
KFU_RE = re.compile(r"^KFU\s*(\d+)", re.I)
NUM_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})\b")
DATE_RE = re.compile(r"(\d{4}-\d{2}-\d{2})")

# AFB.32 lists the tender documents by number; AFB.33 makes later
# supplementary documents (the KFUs) binding.  Lower rank wins a conflict.
CATEGORIES = [
    (lambda n, t: bool(KFU_RE.match(t)) or "ändrings pm" in t.lower(), "Ändrings-PM (KFU)", 1),
    (lambda n, t: "komplettering" in t.lower() or "förtydligand" in t.lower(), "Kompletteringar", 2),
    (lambda n, t: n == "9.1" or re.search(r"\bAF\b", t), "Administrativa föreskrifter", 3),
    (lambda n, t: "avsteg" in t.lower(), "Avsteg mät- och ersättningsregler", 4),
    (lambda n, t: "mängdbeskrivning" in t.lower() or "mängdförteckning" in t.lower(), "Mängdbeskrivning", 5),
    (lambda n, t: "anbudsformulär" in t.lower(), "Anbudsformulär", 6),
    (lambda n, t: "ritningsförteckning" in t.lower(), "Ritningsförteckning", 7),
]


def parse_meta(display_name: str) -> dict:
    """Pull document number, title, revision date and category from a filename."""
    stem = Path(display_name).stem
    number, kfu = None, KFU_RE.match(stem)
    if kfu:
        number = f"KFU {int(kfu.group(1))}"
    else:
        m = NUM_RE.match(stem)
        if m:
            number = f"{int(m.group(1))}.{int(m.group(2))}"

    rev = REV_RE.search(stem)
    rev_date = rev.group(1) if rev else None
    doc_date = rev_date or (DATE_RE.search(stem).group(1) if DATE_RE.search(stem) else None)

    title = stem
    if not kfu and number:
        title = NUM_RE.sub("", title, count=1)
    title = REV_RE.sub("", title)
    title = title.strip(" -–_.")

    category, precedence = "Övrig handling", 8
    for test, name, rank in CATEGORIES:
        if test(number or "", stem):
            category, precedence = name, rank
            break

    return {
        "number": number,
        "title": title or stem,
        "rev_date": rev_date,
        "doc_date": doc_date,
        "category": category,
        "precedence": precedence,
    }


# --------------------------------------------------------------------------
# change marks
# --------------------------------------------------------------------------


def _is_yellow(fill) -> bool:
    return bool(fill) and fill[0] > 0.85 and fill[1] > 0.8 and fill[2] < 0.45


def _is_red(color: int) -> bool:
    r, g, b = (color >> 16) & 255, (color >> 8) & 255, color & 255
    return r > 140 and g < 90 and b < 90


def page_marks(page) -> tuple[set, set]:
    """Return (struck word rects, highlighted word rects) for one page.

    A strike is a thin filled rect or line whose vertical centre falls inside
    the middle band of a word.  Requiring the middle band is what keeps table
    borders and underlines out: those sit at the top or bottom edge of the
    text, not through it.
    """
    words = page.get_text("words")
    struck, highlighted = set(), set()
    page_area = abs(page.rect.get_area()) or 1

    for drawing in page.get_drawings():
        rect = drawing["rect"]
        if rect.is_empty or rect.is_infinite:
            continue

        if _is_yellow(drawing.get("fill")) and abs(rect.get_area()) < page_area * 0.5:
            for w in words:
                wr = pymupdf.Rect(w[:4])
                if wr.intersects(rect):
                    overlap = (wr & rect).get_area()
                    if overlap > wr.get_area() * 0.4:
                        highlighted.add(w[:5])
            continue

        # A table rule runs the width of the page; a strike runs the width of
        # the words it deletes.  Anything wider than most of the page is ruling.
        if rect.height <= 2.5 and 8 <= rect.width <= page.rect.width * 0.7:
            y = (rect.y0 + rect.y1) / 2
            for w in words:
                x0, y0, x1, y1 = w[:4]
                height = y1 - y0
                if not (y0 + height * 0.25 < y < y1 - height * 0.2):
                    continue
                overlap = min(x1, rect.x1) - max(x0, rect.x0)
                if overlap > (x1 - x0) * 0.5:
                    struck.add(w[:5])

    return struck, highlighted


def red_words(page) -> set:
    """Words drawn in red — another common way of marking a change."""
    out = set()
    data = page.get_text("dict")
    for block in data.get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                if not _is_red(span.get("color", 0)):
                    continue
                for w in page.get_text("words"):
                    if pymupdf.Rect(w[:4]).intersects(pymupdf.Rect(span["bbox"])):
                        out.add(w[:5])
    return out


def extract_pdf(path: Path, mark_changes: bool = False) -> tuple[list[str], list[dict]]:
    """Extract page text with change markers, plus a reviewable list of marks.

    ``mark_changes`` is on only for documents that carry a revision date.
    Highlighting is used for ordinary purposes too — the Arbetsmiljöplan
    highlights blank fields to be filled in — so reading yellow as "changed"
    in an unrevised document would be wrong.
    """
    doc = pymupdf.open(path)
    pages, marks = [], []

    for index, page in enumerate(doc, start=1):
        words = page.get_text("words")
        struck, highlighted = page_marks(page) if mark_changes else (set(), set())
        red = red_words(page) if mark_changes else set()

        lines: dict = {}
        for w in words:
            lines.setdefault((w[5], w[6]), []).append(w)

        text_lines, page_marks_found = [], {}
        for key in sorted(lines):
            group = sorted(lines[key], key=lambda w: w[0])
            text = " ".join(w[4] for w in group)
            kinds = set()
            if any(w[:5] in struck for w in group):
                kinds.add("struck")
            if any(w[:5] in highlighted for w in group):
                kinds.add("highlight")
            if any(w[:5] in red for w in group):
                kinds.add("red")

            if "struck" in kinds:
                text = f"[STRUKEN — utgår enligt revidering] {text}"
            elif kinds:
                text = f"[ÄNDRAD i revidering] {text}"
            text_lines.append(text)

            for kind in kinds:
                page_marks_found.setdefault(kind, []).append(" ".join(w[4] for w in group))

        pages.append("\n".join(text_lines))
        for kind, texts in page_marks_found.items():
            for text in texts:
                if text.strip():
                    marks.append({"page": index, "kind": kind, "text": text[:300]})

    doc.close()
    return pages, marks


def extract_xlsx(path: Path) -> tuple[list[str], list[dict]]:
    """One 'page' per worksheet, so citations can point at a sheet."""
    book = load_workbook(path, read_only=True, data_only=True)
    pages = []
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                rows.append(" | ".join(cells))
            if len(rows) > 2000:
                break
        pages.append(f"[Flik: {sheet.title}]\n" + "\n".join(rows))
    book.close()
    return pages, []


# --------------------------------------------------------------------------
# chunking
# --------------------------------------------------------------------------

SECTION_RE = re.compile(r"^\s*(?:\[[^\]]+\]\s*)?([A-ZÅÄÖ]{1,4}\.\d{1,4}(?:\.\d{1,3})?)\b")


def chunk_page(text: str) -> list[tuple[str, str]]:
    """Split a page into chunks, breaking at AMA/AF codes where possible."""
    chunks, buffer, section, buffer_section = [], [], "", ""
    for line in text.split("\n"):
        match = SECTION_RE.match(line)
        if match:
            section = match.group(1)
            if sum(len(b) for b in buffer) > 400:
                chunks.append((buffer_section, "\n".join(buffer)))
                buffer, buffer_section = [], section
        if not buffer:
            buffer_section = section
        buffer.append(line)
        if sum(len(b) for b in buffer) > CHUNK_CHARS:
            chunks.append((buffer_section, "\n".join(buffer)))
            buffer, buffer_section = [], section
    if buffer:
        chunks.append((buffer_section, "\n".join(buffer)))
    return [(s, t) for s, t in chunks if t.strip()]


# --------------------------------------------------------------------------
# KFU change lists
# --------------------------------------------------------------------------

DRAWING_RE = re.compile(r"\b([A-ZÅÄÖ]{1,3}-\d{2}\.\d-\d{4})\b")
HANDLING_RE = re.compile(r"^\s*(\d{1,2}\.\d{1,2})[\.\s]")


def parse_kfu(pages: list[str]) -> list[dict]:
    """Read the 'Handling / Ändringen avser' table of a KFU memo."""
    changes = []
    for text in pages:
        for line in text.split("\n"):
            line = line.strip()
            if not line:
                continue
            drawing = DRAWING_RE.search(line)
            handling = HANDLING_RE.match(line)
            description = DATE_RE.sub("", line).strip()
            if drawing:
                changes.append({"target": drawing.group(1), "kind": "ritning", "description": description})
            elif handling:
                changes.append({"target": handling.group(1), "kind": "handling", "description": description})
    return changes


# --------------------------------------------------------------------------
# database
# --------------------------------------------------------------------------

SCHEMA = """
CREATE TABLE IF NOT EXISTS documents(
    id INTEGER PRIMARY KEY, filename TEXT, display_name TEXT, ext TEXT,
    number TEXT, title TEXT, category TEXT, precedence INTEGER,
    doc_date TEXT, rev_date TEXT, superseded_by INTEGER, page_count INTEGER);
CREATE TABLE IF NOT EXISTS pages(
    doc_id INTEGER, page INTEGER, text TEXT, PRIMARY KEY(doc_id, page));
CREATE TABLE IF NOT EXISTS marks(
    id INTEGER PRIMARY KEY, doc_id INTEGER, page INTEGER, kind TEXT, text TEXT);
CREATE TABLE IF NOT EXISTS kfu_changes(
    id INTEGER PRIMARY KEY, kfu_id INTEGER, target TEXT, kind TEXT,
    description TEXT, target_doc_id INTEGER);
CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5(
    title, section, text, doc_id UNINDEXED, page UNINDEXED,
    tokenize='unicode61 remove_diacritics 2');
"""


def connect(db_path: Path) -> sqlite3.Connection:
    db = sqlite3.connect(db_path, check_same_thread=False)
    db.row_factory = sqlite3.Row
    db.executescript(SCHEMA)
    return db


def reset(db: sqlite3.Connection) -> None:
    for table in ("documents", "pages", "marks", "kfu_changes", "chunks"):
        db.execute(f"DELETE FROM {table}")
    db.commit()


def unpack_zip(zip_path: Path, data_dir: Path) -> int:
    """Extract an FFU zip into the data directory, repairing filenames."""
    count = 0
    with zipfile.ZipFile(zip_path) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            name = fix_name(Path(info.filename).name)
            if name.startswith(("._", "__MACOSX")) or name in SKIP_FILES:
                continue
            if Path(name).suffix.lower() not in SUPPORTED:
                continue
            (data_dir / name).write_bytes(archive.read(info))
            count += 1
    return count


def ingest(db: sqlite3.Connection, data_dir: Path, progress=None) -> dict:
    """Read every document in data_dir into the database. Returns a summary."""
    reset(db)
    paths = sorted(p for p in data_dir.rglob("*") if p.is_file())
    documents = []

    for path in paths:
        name = fix_name(path.name)
        if name in SKIP_FILES or path.suffix.lower() not in SUPPORTED:
            continue
        meta = parse_meta(name)
        try:
            if path.suffix.lower() == ".pdf":
                pages, marks = extract_pdf(path, mark_changes=bool(meta["rev_date"]))
            else:
                pages, marks = extract_xlsx(path)
        except Exception as exc:  # a broken file should not stop the package
            logger.warning("could not read %s: %s", name, exc)
            continue

        cursor = db.execute(
            """INSERT INTO documents(filename, display_name, ext, number, title, category,
                                     precedence, doc_date, rev_date, page_count)
               VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (path.name, name, path.suffix.lower().lstrip("."), meta["number"], meta["title"],
             meta["category"], meta["precedence"], meta["doc_date"], meta["rev_date"], len(pages)),
        )
        doc_id = cursor.lastrowid

        for number, text in enumerate(pages, start=1):
            db.execute("INSERT INTO pages(doc_id, page, text) VALUES(?,?,?)", (doc_id, number, text))
            for section, chunk in chunk_page(text):
                db.execute(
                    "INSERT INTO chunks(title, section, text, doc_id, page) VALUES(?,?,?,?,?)",
                    (f"{meta['number'] or ''} {meta['title']}", section, chunk, doc_id, number),
                )
        for mark in marks:
            db.execute(
                "INSERT INTO marks(doc_id, page, kind, text) VALUES(?,?,?,?)",
                (doc_id, mark["page"], mark["kind"], mark["text"]),
            )

        documents.append({"id": doc_id, "meta": meta, "pages": pages, "marks": len(marks)})
        db.commit()
        if progress:
            progress(f"{name} — {len(pages)} sidor, {len(marks)} ändringsmarkeringar")

    link_revisions(db, documents)
    link_kfu(db, documents)
    db.commit()

    return {
        "documents": len(documents),
        "pages": db.execute("SELECT COUNT(*) FROM pages").fetchone()[0],
        "marks": db.execute("SELECT COUNT(*) FROM marks").fetchone()[0],
        "superseded": db.execute("SELECT COUNT(*) FROM documents WHERE superseded_by IS NOT NULL").fetchone()[0],
    }


def link_revisions(db: sqlite3.Connection, documents: list[dict]) -> None:
    """Point every older version at the newest one with the same number."""
    groups: dict = {}
    for doc in documents:
        number = doc["meta"]["number"]
        if number:
            groups.setdefault(number, []).append(doc)
    for versions in groups.values():
        if len(versions) < 2:
            continue
        versions.sort(key=lambda d: d["meta"]["rev_date"] or d["meta"]["doc_date"] or "")
        newest = versions[-1]
        for old in versions[:-1]:
            db.execute("UPDATE documents SET superseded_by=? WHERE id=?", (newest["id"], old["id"]))


def link_kfu(db: sqlite3.Connection, documents: list[dict]) -> None:
    """Record what each KFU changes, resolving handling numbers to documents."""
    current = db.execute(
        "SELECT id, number, title FROM documents WHERE superseded_by IS NULL"
    ).fetchall()
    by_number = {row["number"]: row["id"] for row in current if row["number"]}

    def resolve(target: str, description: str):
        """KFU 1 names '6.1 Avsteg MER Anläggning 20', the package holds 6.3
        Avsteg MER Anläggning 23.  Fall back to matching on the title so the
        mismatch is visible instead of silently unlinked."""
        if target in by_number:
            return by_number[target]
        words = {w.lower() for w in re.findall(r"[A-Za-zÅÄÖåäö]{4,}", description)}
        best, score = None, 0
        for row in current:
            overlap = len(words & {w.lower() for w in re.findall(r"[A-Za-zÅÄÖåäö]{4,}", row["title"])})
            if overlap > score:
                best, score = row["id"], overlap
        return best if score >= 2 else None
    for doc in documents:
        if doc["meta"]["category"] != "Ändrings-PM (KFU)":
            continue
        for change in parse_kfu(doc["pages"]):
            db.execute(
                "INSERT INTO kfu_changes(kfu_id, target, kind, description, target_doc_id) VALUES(?,?,?,?,?)",
                (doc["id"], change["target"], change["kind"], change["description"],
                 resolve(change["target"], change["description"])),
            )


if __name__ == "__main__":
    import sys

    logging.basicConfig(level=logging.INFO, format="%(message)s")
    base = Path(__file__).resolve().parent
    db = connect(base / "ffu.db")
    summary = ingest(db, Path(sys.argv[1] if len(sys.argv) > 1 else base / "data"), progress=print)
    print(summary)